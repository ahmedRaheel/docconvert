"""PowerPoint -> PDF without LibreOffice. Resolves formatting through the same chain PowerPoint uses
(run -> paragraph -> shape list style -> layout placeholder -> master placeholder -> master text styles ->
presentation defaults) so font size, bold/italic/underline, colour, alignment, bullets, spacing, insets,
autofit scale and vertical anchoring match the slide. Also draws fills, outlines, shape geometry, rotation,
groups, cropped pictures, tables and slide backgrounds (theme colours + lumMod/tint resolved)."""
from __future__ import annotations

import colorsys
import io
import re

from .builtin_fonts import esc, face, families

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS = {"a": A[1:-1], "p": P[1:-1], "r": R[1:-1]}
EMU = 12700.0
_PRESET = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF", "yellow": "FFFF00",
           "gray": "808080", "grey": "808080", "orange": "FFA500"}


class Theme:
    def __init__(self, prs):
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree
        self.colors = {}
        self.map = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
        try:
            master = prs.slide_masters[0]
            blob = master.part.part_related_by(RT.THEME).blob
            root = etree.fromstring(blob)
            for el in root.iterfind(".//a:clrScheme/*", NS):
                name = etree.QName(el).localname
                ch = el[0]
                self.colors[name] = ch.get("lastClr") or ch.get("val") or "000000"
            cm = master._element.find("p:clrMap", NS)
            if cm is not None:
                self.map.update({k: v for k, v in cm.attrib.items()})
        except Exception:
            pass
        self.colors.setdefault("dk1", "000000")
        self.colors.setdefault("lt1", "FFFFFF")
        for i, d in enumerate(["4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"], 1):
            self.colors.setdefault(f"accent{i}", d)
        self.colors.setdefault("hlink", "0563C1")

    def scheme(self, name):
        return self.colors.get(self.map.get(name, name), "000000")

    def color(self, el):
        """a:solidFill / color element -> RRGGBB (applies lumMod, lumOff, tint, shade)."""
        if el is None:
            return None
        if el.tag == A + "solidFill" or el.tag.endswith("Fill") or len(el) and el[0].tag.split("}")[1] in ("srgbClr", "schemeClr", "sysClr", "prstClr"):
            el = el[0] if el.tag.split("}")[1] not in ("srgbClr", "schemeClr", "sysClr", "prstClr") and len(el) else el
        kind = el.tag.split("}")[1]
        if kind == "srgbClr":
            base = el.get("val")
        elif kind == "schemeClr":
            base = self.scheme(el.get("val"))
        elif kind == "sysClr":
            base = el.get("lastClr") or "000000"
        elif kind == "prstClr":
            base = _PRESET.get(el.get("val"), "000000")
        else:
            return None
        r, g, b = (int(base[i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        for m in el:
            t, v = m.tag.split("}")[1], m.get("val")
            if v is None:
                continue
            v = int(v) / 100000
            if t == "lumMod":
                l *= v
            elif t == "lumOff":
                l += v
            elif t == "tint":
                r, g, b = colorsys.hls_to_rgb(h, l, s)
                r, g, b = (c * v + (1 - v) for c in (r, g, b))
                h, l, s = colorsys.rgb_to_hls(r, g, b)
            elif t == "shade":
                l *= v
        r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
        return "%02X%02X%02X" % tuple(round(c * 255) for c in (r, g, b))


def _fill_of(theme, parent):
    """Colour of the first fill child of an spPr/tcPr-like element, 'none' for noFill, None if unspecified."""
    if parent is None:
        return None
    for ch in parent:
        t = ch.tag.split("}")[1]
        if t == "noFill":
            return "none"
        if t == "solidFill":
            return theme.color(ch[0]) if len(ch) else None
        if t == "gradFill":
            gs = ch.find(".//a:gs", NS)
            return theme.color(gs[0]) if gs is not None and len(gs) else None
    return None


class Deck:
    def __init__(self, prs):
        self.prs, self.theme = prs, Theme(prs)
        el = prs.part._element
        self.default_style = el.find("p:defaultTextStyle", NS)
        self.master = prs.slide_masters[0]
        self.tx_styles = self.master._element.find("p:txStyles", NS)

    # ---- inheritance chain of "lvlNpPr" elements for a shape
    def chain(self, slide, shape, lvl):
        tag = f"a:lvl{lvl + 1}pPr"
        out = []
        el = shape._element
        tx = el.find("p:txBody", NS)
        if tx is not None and tx.find("a:lstStyle", NS) is not None:
            out.append(tx.find(f"a:lstStyle/{tag}", NS))
        kind = "other"
        try:
            if shape.is_placeholder:
                ph = shape.placeholder_format
                tname = ph.type.name if ph.type is not None else "OBJECT"
                kind = "title" if tname in ("TITLE", "CENTER_TITLE") else "body" if tname in ("BODY", "OBJECT", "SUBTITLE", "VERTICAL_BODY", "TABLE", "CHART") else "other"
                for holder in (slide.slide_layout, self.master):
                    for cand in holder.placeholders:
                        same = cand.placeholder_format.idx == ph.idx or (kind == "title" and (cand.placeholder_format.type is not None and cand.placeholder_format.type.name in ("TITLE", "CENTER_TITLE"))) \
                            or (kind == "body" and holder is self.master and cand.placeholder_format.type is not None and cand.placeholder_format.type.name == "BODY")
                        if same:
                            ctx = cand._element.find("p:txBody", NS)
                            if ctx is not None:
                                out.append(ctx.find(f"a:lstStyle/{tag}", NS))
                            break
        except Exception:
            pass
        if self.tx_styles is not None:
            style = {"title": "titleStyle", "body": "bodyStyle"}.get(kind, "otherStyle")
            out.append(self.tx_styles.find(f"p:{style}/{tag}", NS))
        if self.default_style is not None:
            out.append(self.default_style.find(tag, NS))
        return [o for o in out if o is not None]

    def body_pr(self, slide, shape):
        """bodyPr attributes with placeholder inheritance."""
        out = {}
        sources = []
        try:
            if shape.is_placeholder:
                for holder in (self.master, slide.slide_layout):
                    for cand in holder.placeholders:
                        if cand.placeholder_format.idx == shape.placeholder_format.idx or (
                                cand.placeholder_format.type == shape.placeholder_format.type):
                            bp = cand._element.find("p:txBody/a:bodyPr", NS)
                            if bp is not None:
                                sources.append(bp)
                            break
        except Exception:
            pass
        bp = shape._element.find("p:txBody/a:bodyPr", NS)
        if bp is not None:
            sources.append(bp)
        for bp in sources:
            out.update(bp.attrib)
            af = bp.find("a:normAutofit", NS)
            if af is not None:
                out["fontScale"] = af.get("fontScale")
        return out


def _first(chain, getter):
    for el in chain:
        v = getter(el)
        if v is not None:
            return v
    return None


def paragraph_markup(deck, slide, shape, p_el, bodyscale=1.0):
    """-> (markup, style kwargs, bullet text) for an a:p element."""
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    theme = deck.theme
    ppr = p_el.find("a:pPr", NS)
    lvl = int(ppr.get("lvl", 0)) if ppr is not None else 0
    chain = ([ppr] if ppr is not None else []) + deck.chain(slide, shape, lvl)
    d_chain = [c.find("a:defRPr", NS) for c in chain if c is not None]
    d_chain = [d for d in d_chain if d is not None]

    algn = _first(chain, lambda e: e.get("algn"))
    marL = _first(chain, lambda e: e.get("marL"))
    indent = _first(chain, lambda e: e.get("indent"))
    lnspc = _first(chain, lambda e: (e.find("a:lnSpc/a:spcPct", NS).get("val") if e.find("a:lnSpc/a:spcPct", NS) is not None else None))
    spc_bef = _first(chain, lambda e: (e.find("a:spcBef/a:spcPts", NS).get("val") if e.find("a:spcBef/a:spcPts", NS) is not None else None))
    spc_aft = _first(chain, lambda e: (e.find("a:spcAft/a:spcPts", NS).get("val") if e.find("a:spcAft/a:spcPts", NS) is not None else None))
    bu = None
    for e in chain:
        if e.find("a:buNone", NS) is not None:
            bu = None
            break
        ch = e.find("a:buChar", NS)
        if ch is not None:
            bu = ch.get("char", "•")
            break
        if e.find("a:buAutoNum", NS) is not None:
            bu = "#"
            break

    parts, max_size = [], 0.0
    end = p_el.find("a:endParaRPr", NS)
    for node in p_el:
        kind = node.tag.split("}")[1]
        if kind == "br":
            parts.append("<br/>")
            continue
        if kind not in ("r", "fld"):
            continue
        rpr = node.find("a:rPr", NS)
        t = node.find("a:t", NS)
        text = t.text if t is not None and t.text else ""
        chain_r = ([rpr] if rpr is not None else []) + d_chain

        def attr(name, chain_r=chain_r):
            return _first(chain_r, lambda e: e.get(name))
        size = float(attr("sz") or 1800) / 100 * bodyscale
        bold = attr("b") in ("1", "true")
        italic = attr("i") in ("1", "true")
        under = (attr("u") or "none") != "none"
        strike = (attr("strike") or "noStrike") != "noStrike"
        base = attr("baseline")
        fill = _first(chain_r, lambda e: e.find("a:solidFill", NS))
        col = theme.color(fill) if fill is not None else None
        if col is None and shape._element.find("p:style/a:fontRef", NS) is not None:
            col = theme.color(shape._element.find("p:style/a:fontRef", NS)[0]) if len(shape._element.find("p:style/a:fontRef", NS)) else None
        typeface = _first(chain_r, lambda e: (e.find("a:latin", NS).get("typeface") if e.find("a:latin", NS) is not None else None))
        max_size = max(max_size, size)
        link = rpr.find("a:hlinkClick", NS) if rpr is not None else None
        if link is not None and not fill:
            col, under = theme.scheme("hlink"), True
        m = esc(text).replace("\x0b", "<br/>")
        attrs = f'name="{face(typeface, bold, italic)}" size="{size:.1f}"' + (f' color="#{col}"' if col else ' color="#000000"')
        m = f"<font {attrs}>{m}</font>"
        if under:
            m = f"<u>{m}</u>"
        if strike:
            m = f"<strike>{m}</strike>"
        if base and base.lstrip("-").isdigit():
            m = f"<super>{m}</super>" if int(base) > 0 else f"<sub>{m}</sub>"
        parts.append(m)
    if not max_size:
        er = end.get("sz") if end is not None else None
        max_size = float(er or _first(d_chain, lambda e: e.get("sz")) or 1800) / 100 * bodyscale
    style = dict(
        alignment={"ctr": TA_CENTER, "r": TA_RIGHT, "just": TA_JUSTIFY}.get(algn, TA_LEFT),
        leftIndent=int(marL or 0) / EMU, firstLineIndent=int(indent or 0) / EMU if bu is None else 0,
        spaceBefore=int(spc_bef or 0) / 100, spaceAfter=int(spc_aft or 0) / 100,
        leading=max_size * 1.2 * (int(lnspc) / 100000 if lnspc else 1.0), fontSize=max_size)
    return "".join(parts), style, bu, int(indent or 0) / EMU


def draw_text(c, deck, slide, shape, txbody, x, y, w, h, scale_font=True, default_insets=(91440, 45720, 91440, 45720), props=None):
    """Draw an a:txBody inside the box (x,y top-left, points, y downward)."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    bp = deck.body_pr(slide, shape) if props is None else props
    l, t, r_, b = (int(bp.get(k, d)) / EMU for k, d in zip(("lIns", "tIns", "rIns", "bIns"), default_insets))
    fs = float(bp["fontScale"]) / 100000 if bp.get("fontScale") else 1.0
    paras, counter = [], 0
    for p_el in txbody.findall("a:p", NS):
        markup, style, bu, hang = paragraph_markup(deck, slide, shape, p_el, fs)
        bullet = None
        if bu == "#":
            counter += 1
            bullet = f"{counter}."
        elif bu:
            counter = 0
            bullet = bu if ord(bu[0]) < 0xF000 and bu not in ("Ø", "§", "ü") else "•"
        else:
            counter = 0
        st = ParagraphStyle("t", fontName=face(None), **{**style, "firstLineIndent": 0 if bullet else style["firstLineIndent"]},
                            autoLeading="max", bulletIndent=max(style["leftIndent"] + hang, 0),
                            bulletFontName=face(None), bulletFontSize=style["fontSize"])
        if bullet and style["leftIndent"] == 0:
            st.leftIndent = 18
        paras.append(Paragraph(markup or "&nbsp;", st, bulletText=bullet))
    avail_w = max(w - l - r_, 10)
    heights = []
    for p in paras:
        _, ph = p.wrap(avail_w, 10 ** 6)
        heights.append(ph + p.style.spaceBefore + p.style.spaceAfter)
    total = sum(heights)
    anchor = bp.get("anchor", "t")
    top = y + t if anchor == "t" else y + (h - total) / 2 if anchor == "ctr" else y + h - total - b
    page_h = c._pagesize[1]
    cur = top
    for p, hh in zip(paras, heights):
        cur += p.style.spaceBefore
        p.drawOn(c, x + l, page_h - cur - (hh - p.style.spaceBefore - p.style.spaceAfter))
        cur += hh - p.style.spaceBefore


def pptx_to_pdf(src, dst):
    from pptx import Presentation
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    families()
    prs = Presentation(str(src))
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    deck = Deck(prs)
    c = canvas.Canvas(str(dst), pagesize=(W, H))
    th = deck.theme

    def bg_color(slide):
        for holder in (slide, slide.slide_layout, deck.master):
            bg = holder._element.find("p:cSld/p:bg", NS)
            if bg is not None:
                fill = _fill_of(th, bg.find("p:bgPr", NS))
                if fill and fill != "none":
                    return fill
                ref = bg.find("p:bgRef", NS)
                if ref is not None and len(ref):
                    return th.color(ref[0])
        return "FFFFFF"

    def hexrgb(hx):
        return tuple(int(hx[i:i + 2], 16) / 255 for i in (0, 2, 4))

    def geometry(shape):
        g = shape._element.find("p:spPr/a:prstGeom", NS)
        return g.get("prst") if g is not None else "rect"

    def draw_shapes(slide, shapes, aff, only_nonplaceholder=False):
        for shape in shapes:
            if only_nonplaceholder and getattr(shape, "is_placeholder", False):
                continue
            el = shape._element
            if el.find(".//p:nvPr", NS) is not None and el.find(".//p:nvPr", NS).get("userDrawn") == "0":
                pass
            try:
                ax, bx, ay, by = aff
                x = ax * shape.left + bx
                y = ay * shape.top + by
                w, h = ax * shape.width, ay * shape.height
            except Exception:
                continue
            x, y, w, h = x / EMU, y / EMU, w / EMU, h / EMU
            if shape.shape_type == 6:  # group
                xf = el.find("p:grpSpPr/a:xfrm", NS)
                ch_off, ch_ext = xf.find("a:chOff", NS), xf.find("a:chExt", NS)
                sx = shape.width / int(ch_ext.get("cx")) if int(ch_ext.get("cx")) else 1
                sy = shape.height / int(ch_ext.get("cy")) if int(ch_ext.get("cy")) else 1
                gx, gy = shape.left - int(ch_off.get("x")) * sx, shape.top - int(ch_off.get("y")) * sy
                new = (ax * sx, ax * gx + bx, ay * sy, ay * gy + by)
                draw_shapes(slide, shape.shapes, new)
                continue
            rot = shape.rotation if hasattr(shape, "rotation") else 0
            c.saveState()
            page_y = H - y - h
            if rot:
                c.translate(x + w / 2, H - (y + h / 2))
                c.rotate(-rot)
                c.translate(-(x + w / 2), -(H - (y + h / 2)))
            sp = el.find("p:spPr", NS)
            fill = _fill_of(th, sp)
            ln = sp.find("a:ln", NS) if sp is not None else None
            line_col = None
            line_w = 0.75
            if ln is not None:
                line_col = _fill_of(th, ln)
                if ln.get("w"):
                    line_w = int(ln.get("w")) / EMU
            style = el.find("p:style", NS)
            if style is not None:
                if fill is None and style.find("a:fillRef", NS) is not None and style.find("a:fillRef", NS).get("idx") not in ("0", None) and len(style.find("a:fillRef", NS)):
                    fill = th.color(style.find("a:fillRef", NS)[0])
                lr = style.find("a:lnRef", NS)
                if line_col is None and lr is not None and lr.get("idx") not in ("0", None) and len(lr):
                    line_col = th.color(lr[0])
            if fill == "none":
                fill = None
            if line_col == "none":
                line_col = None
            prst = geometry(shape)
            is_line = el.tag == P + "cxnSp" or prst in ("line", "straightConnector1")
            if is_line:
                flip_h = el.find("p:spPr/a:xfrm", NS) is not None and el.find("p:spPr/a:xfrm", NS).get("flipH") == "1"
                flip_v = el.find("p:spPr/a:xfrm", NS) is not None and el.find("p:spPr/a:xfrm", NS).get("flipV") == "1"
                c.setStrokeColorRGB(*hexrgb(line_col or "000000"))
                c.setLineWidth(line_w)
                x1, x2 = (x + w, x) if flip_h else (x, x + w)
                y1, y2 = (y + h, y) if flip_v else (y, y + h)
                c.line(x1, H - y1, x2, H - y2)
            elif shape.shape_type == 13:     # picture
                try:
                    from PIL import Image
                    im = Image.open(io.BytesIO(shape.image.blob))
                    cl, cr, ct, cb = shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom
                    if any((cl, cr, ct, cb)):
                        iw, ih = im.size
                        im = im.crop((int(cl * iw), int(ct * ih), int(iw - cr * iw), int(ih - cb * ih)))
                    c.drawImage(ImageReader(im), x, page_y, w, h, mask="auto")
                except Exception:
                    pass
            elif getattr(shape, "has_table", False) and shape.has_table:
                draw_table(slide, shape, x, y, w, h)
            else:
                if fill or line_col:
                    if fill:
                        c.setFillColorRGB(*hexrgb(fill))
                    if line_col:
                        c.setStrokeColorRGB(*hexrgb(line_col))
                        c.setLineWidth(line_w)
                    kw = dict(stroke=1 if line_col else 0, fill=1 if fill else 0)
                    if prst == "ellipse":
                        c.ellipse(x, page_y, x + w, page_y + h, **kw)
                    elif prst in ("roundRect", "round2SameRect", "snipRoundRect"):
                        c.roundRect(x, page_y, w, h, min(w, h) * 0.16, **kw)
                    elif prst in ("triangle", "rtTriangle"):
                        pth = c.beginPath()
                        pth.moveTo(x + w / 2, page_y + h)
                        pth.lineTo(x, page_y)
                        pth.lineTo(x + w, page_y)
                        pth.close()
                        c.drawPath(pth, **kw)
                    elif prst == "diamond":
                        pth = c.beginPath()
                        pth.moveTo(x + w / 2, page_y + h)
                        pth.lineTo(x + w, page_y + h / 2)
                        pth.lineTo(x + w / 2, page_y)
                        pth.lineTo(x, page_y + h / 2)
                        pth.close()
                        c.drawPath(pth, **kw)
                    else:
                        c.rect(x, page_y, w, h, **kw)
                tx = el.find("p:txBody", NS)
                if tx is not None and any((p.find("a:r", NS) is not None or p.find("a:fld", NS) is not None) for p in tx.findall("a:p", NS)):
                    draw_text(c, deck, slide, shape, tx, x, y, w, h)
            c.restoreState()

    def draw_table(slide, shape, x, y, w, h):
        from reportlab.lib.styles import ParagraphStyle
        tbl = shape.table
        tbl_el = tbl._tbl
        col_w = [col.width / EMU for col in tbl.columns]
        row_h = [row.height / EMU for row in tbl.rows]
        header = tbl.first_row
        banded = tbl.horz_banding
        accent = th.scheme("accent1")
        cy = y
        for ri, row in enumerate(tbl_el.findall("a:tr", NS)):
            cx = x
            tcs = row.findall("a:tc", NS)
            for ci, tc in enumerate(tcs):
                cw = col_w[ci] if ci < len(col_w) else 60
                if tc.get("hMerge") == "1" or tc.get("vMerge") == "1":
                    cx += cw
                    continue
                span_w = sum(col_w[ci:ci + int(tc.get("gridSpan", 1))])
                span_h = sum(row_h[ri:ri + int(tc.get("rowSpan", 1))])
                tcpr = tc.find("a:tcPr", NS)
                fill = _fill_of(th, tcpr)
                if fill is None and (header or banded):
                    if header and ri == 0:
                        fill = accent
                    elif banded:
                        fill = th.color(_mk_tint(accent, 0.4 if ri % 2 else 0.2))
                if fill and fill != "none":
                    c.setFillColorRGB(*hexrgb(fill))
                    c.rect(cx, H - cy - span_h, span_w, span_h, stroke=0, fill=1)
                for tag, sx0, sy0, sx1, sy1 in (("lnT", cx, cy, cx + span_w, cy), ("lnB", cx, cy + span_h, cx + span_w, cy + span_h),
                                                ("lnL", cx, cy, cx, cy + span_h), ("lnR", cx + span_w, cy, cx + span_w, cy + span_h)):
                    ln = tcpr.find(f"a:{tag}", NS) if tcpr is not None else None
                    col = _fill_of(th, ln) if ln is not None else None
                    if col and col != "none":
                        c.setStrokeColorRGB(*hexrgb(col))
                        c.setLineWidth(int(ln.get("w", 12700)) / EMU)
                        c.line(sx0, H - sy0, sx1, H - sy1)
                tx = tc.find("a:txBody", NS)
                if tx is not None:
                    if header and ri == 0 and fill == accent:           # table style: white bold header text
                        for r_el in tx.iter(A + "r"):
                            rpr = r_el.find("a:rPr", NS)
                            if rpr is None:
                                from lxml import etree
                                rpr = etree.SubElement(r_el, A + "rPr")
                                r_el.insert(0, rpr)
                            rpr.set("b", "1")
                            if rpr.find("a:solidFill", NS) is None:
                                from lxml import etree
                                sf = etree.SubElement(rpr, A + "solidFill")
                                etree.SubElement(sf, A + "srgbClr").set("val", "FFFFFF")
                    props = {k: tcpr.get(k) for k in ("marL", "marT", "marR", "marB") if tcpr is not None and tcpr.get(k)}
                    bp = {"lIns": props.get("marL", 91440), "tIns": props.get("marT", 45720), "rIns": props.get("marR", 91440),
                          "bIns": props.get("marB", 45720), "anchor": {"ctr": "ctr", "b": "b"}.get(tcpr.get("anchor") if tcpr is not None else None, "t")}
                    draw_text(c, deck, slide, shape, tx, cx, cy, span_w, span_h, props=bp)
                cx += cw
            cy += row_h[ri]

    for slide in prs.slides:
        c.setFillColorRGB(*hexrgb(bg_color(slide)))
        c.rect(0, 0, W, H, stroke=0, fill=1)
        ident = (1.0, 0.0, 1.0, 0.0)
        draw_shapes(slide, deck.master.shapes, ident, only_nonplaceholder=True)
        draw_shapes(slide, slide.slide_layout.shapes, ident, only_nonplaceholder=True)
        draw_shapes(slide, slide.shapes, ident)
        c.showPage()
    if not len(prs.slides):
        c.showPage()
    c.save()


def _mk_tint(hexcolor, amount):
    from lxml import etree
    el = etree.Element(A + "srgbClr", val=hexcolor)
    etree.SubElement(el, A + "tint", val=str(int(amount * 100000)))
    return el
