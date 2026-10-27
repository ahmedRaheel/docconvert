"""Word / Excel / PowerPoint -> PDF.

Engines
-------
``libreoffice``  highest fidelity; needs LibreOffice installed (``soffice`` on PATH or
                 ``DOCCONVERT_SOFFICE=/path/to/soffice``).
``builtin``      pure Python (reportlab). Good for text, tables, pictures and basic
                 styling; no complex layout (floating objects, charts, SmartArt).
``auto``         libreoffice if found, otherwise builtin.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from .errors import ConversionError, MissingDependencyError

_SOFFICE_PATHS = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]


def find_soffice() -> str | None:
    env = os.environ.get("DOCCONVERT_SOFFICE")
    if env and Path(env).exists():
        return env
    return shutil.which("soffice") or shutil.which("libreoffice") or next(
        (p for p in _SOFFICE_PATHS if Path(p).exists()), None)


def office_to_pdf(src, dst, *, engine: str = "auto", timeout: int = 180, **_):
    if engine not in ("auto", "libreoffice", "builtin"):
        engine = "auto"
    exe = find_soffice() if engine != "builtin" else None
    if engine == "libreoffice" and not exe:
        raise MissingDependencyError("LibreOffice not found. Install it or set DOCCONVERT_SOFFICE.")
    if exe:
        try:
            return _libreoffice(exe, src, dst, timeout)
        except Exception:
            if engine == "libreoffice":
                raise
    return _builtin(src, dst)


def _libreoffice(exe, src, dst, timeout, filter_name="pdf", copy=True):
    src = Path(src)
    with tempfile.TemporaryDirectory(prefix="docconvert_") as td:
        td = Path(td)
        if copy:  # avoids odd characters in file names
            tmp_in = td / ("input" + src.suffix.lower())
            shutil.copyfile(src, tmp_in)
        else:  # keep original location so relative images/CSS still resolve
            tmp_in = src.resolve()
        out = td / "out"
        out.mkdir()
        cmd = [exe, f"-env:UserInstallation={(td / 'profile').as_uri()}", "--headless", "--norestore",
               "--convert-to", filter_name, "--outdir", str(out), str(tmp_in)]
        env = os.environ.copy()
        env.setdefault("SAL_USE_VCLPLUGIN", "svp")
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=env)
        except subprocess.TimeoutExpired as exc:
            raise ConversionError(f"LibreOffice timed out after {timeout}s") from exc
        produced = out / (tmp_in.stem + ".pdf")
        if not produced.exists():
            raise ConversionError(f"LibreOffice produced no PDF: {res.stderr.strip() or res.stdout.strip()}")
        shutil.move(str(produced), str(dst))
    return dst


# --------------------------------------------------------------------------- builtin
def _fonts():
    """Register a Unicode-capable font family when one is available; else Helvetica."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    sets = [
        ("/usr/share/fonts/truetype/dejavu/", ["DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"]),
        ("/usr/share/fonts/truetype/liberation/", ["LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf"]),
        (r"C:\Windows\Fonts\\", ["arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"]),
        ("/Library/Fonts/", ["Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf"]),
    ]
    for folder, files in sets:
        paths = [Path(folder) / f for f in files]
        if all(p.exists() for p in paths):
            names = ["DC", "DC-B", "DC-I", "DC-BI"]
            for n, p in zip(names, paths):
                if n not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(n, str(p)))
            for b, i, n in ((0, 0, names[0]), (1, 0, names[1]), (0, 1, names[2]), (1, 1, names[3])):
                addMapping("DC", b, i, n)
            return tuple(names)
    return ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique")


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _builtin(src, dst):
    ext = Path(src).suffix.lower()
    fn = {".docx": _docx_pdf, ".xlsx": _xlsx_pdf, ".xlsm": _xlsx_pdf, ".pptx": _pptx_pdf}.get(ext)
    if fn is None:
        raise MissingDependencyError(f"Built-in engine cannot read {ext}; install LibreOffice.")
    fn(Path(src), Path(dst), _fonts())
    return dst


def _docx_pdf(src, dst, fonts):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph as DPara
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image as RImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from reportlab.platypus import Table as RTable, TableStyle

    doc = Document(str(src))
    sec = doc.sections[0]
    pw, ph = sec.page_width.pt, sec.page_height.pt
    lm, rm, tm, bm = (sec.left_margin.pt, sec.right_margin.pt, sec.top_margin.pt, sec.bottom_margin.pt)
    avail = pw - lm - rm
    base = ParagraphStyle("b", fontName=fonts[0], fontSize=11, leading=14, spaceAfter=6)
    sizes = {"Title": 24, "Heading 1": 18, "Heading 2": 15, "Heading 3": 13, "Heading 4": 12}
    aligns = {0: TA_LEFT, 1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}

    def runs_markup(p):
        out = []
        for r in p.runs:
            t = _esc(r.text)
            if not t:
                continue
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            if r.underline:
                t = f"<u>{t}</u>"
            try:
                if r.font.color is not None and r.font.color.rgb is not None:
                    t = f'<font color="#{r.font.color.rgb}">{t}</font>'
            except Exception:
                pass
            out.append(t)
        return "".join(out)

    def para_flowables(p):
        flow = []
        for r in p.runs:
            for rid in r._r.xpath(".//a:blip/@r:embed"):
                try:
                    blob = doc.part.related_parts[rid].blob
                    iw, ih = ImageReader(io.BytesIO(blob)).getSize()
                    w = min(iw * 0.75, avail)
                    flow.append(RImage(io.BytesIO(blob), width=w, height=w * ih / iw))
                except Exception:
                    pass
            if r._r.xpath('.//w:br[@w:type="page"]'):
                flow.append(PageBreak())
        text = runs_markup(p)
        name = p.style.name if p.style is not None else "Normal"
        if text.strip():
            st = ParagraphStyle(name, parent=base, alignment=aligns.get(int(p.alignment or 0), TA_LEFT))
            if name in sizes:
                st.fontName, st.fontSize, st.leading, st.spaceBefore = fonts[1], sizes[name], sizes[name] * 1.25, 8
                text = _esc(p.text)
            bullet = None
            if name.startswith("List Bullet"):
                bullet, st.leftIndent, st.bulletIndent = "•", 18, 6
            elif name.startswith("List Number"):
                bullet, st.leftIndent, st.bulletIndent = "–", 18, 6
            flow.insert(0, Paragraph(text, st, bulletText=bullet)) if not flow else flow.append(Paragraph(text, st, bulletText=bullet))
        elif not flow:
            flow.append(Spacer(1, 8))
        return flow

    story = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            story += para_flowables(DPara(child, doc))
        elif child.tag == qn("w:tbl"):
            t = Table(child, doc)
            cell_style = ParagraphStyle("c", parent=base, fontSize=9, leading=11, spaceAfter=0)
            data = [[Paragraph(_esc(c.text).replace("\n", "<br/>"), cell_style) for c in row.cells] for row in t.rows]
            if data:
                rt = RTable(data, repeatRows=1)
                rt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [rt, Spacer(1, 8)]
    if not story:
        story = [Spacer(1, 1)]
    SimpleDocTemplate(str(dst), pagesize=(pw, ph), leftMargin=lm, rightMargin=rm, topMargin=tm, bottomMargin=bm).build(story)


def _xlsx_pdf(src, dst, fonts):
    import datetime as dt

    from openpyxl import load_workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from reportlab.platypus import Table as RTable, TableStyle

    wb_v = load_workbook(str(src), data_only=True)
    wb_f = load_workbook(str(src), data_only=False)
    page = landscape(A4)
    avail = page[0] - 60
    story = []
    cs = ParagraphStyle("c", fontName=fonts[0], fontSize=8, leading=10)
    title = ParagraphStyle("t", fontName=fonts[1], fontSize=13, spaceAfter=6)

    def fmt(v):
        if isinstance(v, float):
            return f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) >= 1000 else f"{v:.10g}"
        if isinstance(v, (dt.datetime, dt.date)):
            return v.strftime("%Y-%m-%d")
        return "" if v is None else str(v)

    for si, ws in enumerate(wb_v.worksheets):
        wf = wb_f[ws.title]
        if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
            continue
        if story:
            story.append(PageBreak())
        story += [Paragraph(_esc(ws.title), title)]
        data, styles = [], [("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for r_i, row in enumerate(ws.iter_rows()):
            out = []
            for c_i, c in enumerate(row):
                v = c.value if c.value is not None else wf.cell(row=c.row, column=c.column).value
                text = _esc(fmt(v))
                if c.font and c.font.b:
                    text = f"<b>{text}</b>"
                out.append(Paragraph(text, cs))
                try:
                    if c.fill and c.fill.fill_type == "solid" and isinstance(c.fill.fgColor.rgb, str) and len(c.fill.fgColor.rgb) == 8:
                        styles.append(("BACKGROUND", (c_i, r_i), (c_i, r_i), colors.HexColor("#" + c.fill.fgColor.rgb[2:])))
                except Exception:
                    pass
            data.append(out)
        ncols = max(len(r) for r in data)
        widths = [(ws.column_dimensions[chr(65 + i)].width or 10) * 5.2 if i < 26 else 52 for i in range(ncols)]
        k = min(1.0, avail / sum(widths))
        t = RTable(data, colWidths=[w * k for w in widths], repeatRows=1)
        t.setStyle(TableStyle(styles))
        story.append(t)
    if not story:
        story = [Spacer(1, 1)]
    SimpleDocTemplate(str(dst), pagesize=page, leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30).build(story)


def _pptx_pdf(src, dst, fonts):
    from pptx import Presentation
    from pptx.util import Pt
    from reportlab.lib.utils import ImageReader, simpleSplit
    from reportlab.pdfgen import canvas

    prs = Presentation(str(src))
    W, H = prs.slide_width.pt, prs.slide_height.pt
    c = canvas.Canvas(str(dst), pagesize=(W, H))

    def rgb(color_fmt):
        try:
            v = color_fmt.rgb
            return (v[0] / 255, v[1] / 255, v[2] / 255) if v is not None else None
        except Exception:
            return None

    def draw(shapes):
        for sh in shapes:
            try:
                x, y, w, h = sh.left.pt, sh.top.pt, sh.width.pt, sh.height.pt
            except Exception:
                continue
            if sh.shape_type == 6:  # group
                draw(sh.shapes)
                continue
            try:
                if sh.fill.type == 1 and rgb(sh.fill.fore_color):
                    c.setFillColorRGB(*rgb(sh.fill.fore_color))
                    c.rect(x, H - y - h, w, h, stroke=0, fill=1)
            except Exception:
                pass
            if sh.shape_type == 13:
                try:
                    c.drawImage(ImageReader(io.BytesIO(sh.image.blob)), x, H - y - h, w, h, mask="auto")
                except Exception:
                    pass
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = sh.table.rows
                rh = h / max(len(rows), 1)
                for r_i, row in enumerate(rows):
                    cw = w / max(len(row.cells), 1)
                    for c_i, cell in enumerate(row.cells):
                        cx, cy = x + c_i * cw, H - y - (r_i + 1) * rh
                        c.setStrokeColorRGB(0.6, 0.6, 0.6)
                        c.rect(cx, cy, cw, rh, stroke=1, fill=0)
                        c.setFillColorRGB(0, 0, 0)
                        c.setFont(fonts[1] if r_i == 0 else fonts[0], 9)
                        c.drawString(cx + 3, cy + rh / 2 - 3, simpleSplit(cell.text, fonts[0], 9, cw - 6)[0] if cell.text else "")
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                cursor = H - y - 4
                for para in sh.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs) or para.text
                    if not text.strip():
                        cursor -= 10
                        continue
                    r0 = para.runs[0] if para.runs else None
                    size = (r0.font.size.pt if r0 is not None and r0.font.size else 18)
                    size = min(size, 60)
                    bold = bool(r0 is not None and r0.font.bold)
                    italic = bool(r0 is not None and r0.font.italic)
                    face = fonts[3] if bold and italic else fonts[1] if bold else fonts[2] if italic else fonts[0]
                    col = rgb(r0.font.color) if r0 is not None else None
                    c.setFillColorRGB(*(col or (0, 0, 0)))
                    for line in simpleSplit(text, face, size, max(w - 8, 20)):
                        cursor -= size * 1.2
                        c.setFont(face, size)
                        c.drawString(x + 4, cursor, line)
                    cursor -= size * 0.3

    for slide in prs.slides:
        try:
            bg = slide.background.fill
            if bg.type == 1 and rgb(bg.fore_color):
                c.setFillColorRGB(*rgb(bg.fore_color))
                c.rect(0, 0, W, H, stroke=0, fill=1)
        except Exception:
            pass
        draw(slide.shapes)
        c.showPage()
    if not len(prs.slides):
        c.showPage()
    c.save()
