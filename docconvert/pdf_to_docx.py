"""PDF -> Word.

Engines
-------
``pdf2docx``  best layout fidelity (needs the optional ``pdf2docx`` package, which uses PyMuPDF/AGPL).
``builtin``   pure pdfplumber + python-docx: paragraphs, headings, bullets, bold/italic/colour runs,
              tables and pictures, in reading order.
``auto``      pdf2docx when installed, otherwise builtin.
"""
from __future__ import annotations

import io
import re

from .errors import ConversionError, MissingDependencyError
from .utils import chars_to_segments, group_runs, most_common_size, parse_pages

_BULLET = re.compile(r"^[•●▪■◦‣\-–—*]\s+")


def pdf_to_docx(src, dst, *, engine: str = "auto", pages=None, layout: str = "auto", **_):
    """layout: ``"auto"`` (exact copy for forms, flowing text otherwise), ``"fixed"`` or ``"flow"``."""
    if engine not in ("auto", "pdf2docx", "builtin"):
        engine = "auto"
    layout = {"exact": "fixed", "flowing": "flow"}.get(layout, layout)
    if layout not in ("auto", "fixed", "flow"):
        layout = "auto"
    if layout != "flow":
        from .pdf_to_docx_fixed import looks_like_form, pdf_to_docx_fixed
        if layout == "fixed" or looks_like_form(src, pages):
            return pdf_to_docx_fixed(src, dst, pages=pages)
    if engine in ("auto", "pdf2docx"):
        try:
            return _with_pdf2docx(src, dst, pages)
        except ImportError:
            if engine == "pdf2docx":
                raise MissingDependencyError("Install the optional dependency: pip install pdf2docx")
        except Exception as exc:  # pdf2docx can choke on unusual PDFs -> fall back
            if engine == "pdf2docx":
                raise ConversionError(f"pdf2docx failed: {exc}") from exc
    return _builtin(src, dst, pages)


def _with_pdf2docx(src, dst, pages):
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from pdf2docx import Converter

    cv = Converter(str(src))
    try:
        if pages is None:
            cv.convert(str(dst))
        else:
            from pypdf import PdfReader

            n = len(PdfReader(str(src)).pages)
            cv.convert(str(dst), pages=parse_pages(pages, n))
    finally:
        cv.close()
    return dst


def _builtin(src, dst, pages):
    import pdfplumber
    import pypdfium2 as pdfium
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    pdoc = pdfium.PdfDocument(str(src))
    doc = Document()
    with pdfplumber.open(str(src)) as pdf:
        idx = parse_pages(pages, len(pdf.pages))
        first = pdf.pages[idx[0]]
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Pt(first.width), Pt(first.height)
        margin = Pt(54)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margin
        avail = float(first.width) - 108

        body = most_common_size(c["size"] for i in idx for c in pdf.pages[i].chars)

        for n, i in enumerate(idx):
            page = pdf.pages[i]
            if n:
                doc.add_page_break()
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]

            def outside(obj):
                cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
                return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes)

            flow = []  # (top, kind, payload)
            for ln in page.filter(outside).extract_text_lines(return_chars=True, strip=True):
                segs = chars_to_segments(ln["chars"], gap_factor=6)
                if not segs:
                    continue
                flow.append((ln["top"], "line", dict(
                    top=ln["top"], bottom=ln["bottom"], x0=ln["x0"], x1=ln["x1"],
                    runs=[r for s in segs for r in group_runs(s) + [dict(text="  ", font=None, bold=False, italic=False, size=0, color="000000")]][:-1],
                    size=max((r["size"] for s in segs for r in group_runs(s)), default=body))))
            for t in tables:
                flow.append((t.bbox[1], "table", t))
            if page.images:
                full = pdoc[i].render(scale=150 / 72).to_pil().convert("RGB")
                for im in page.images:
                    if im["x1"] - im["x0"] > 4 and im["bottom"] - im["top"] > 4:
                        flow.append((im["top"], "image", (im, full)))
            flow.sort(key=lambda t: t[0])

            # group consecutive lines into paragraphs
            blocks, cur = [], None
            for _, kind, payload in flow:
                if kind != "line":
                    cur = None
                    blocks.append((kind, payload))
                    continue
                p = payload
                if cur and (p["top"] - cur["bottom"] <= 0.7 * cur["size"]
                            and abs(p["size"] - cur["size"]) < 1.5 and p["x0"] - cur["x0"] < 14
                            and not _BULLET.match(p["runs"][0]["text"])):
                    if cur["runs"] and not cur["runs"][-1]["text"].endswith(" "):
                        cur["runs"][-1]["text"] += " "
                    cur["runs"] += p["runs"]
                    cur["bottom"], cur["x1"] = p["bottom"], max(cur["x1"], p["x1"])
                else:
                    cur = dict(p)
                    blocks.append(("para", cur))

            for kind, p in blocks:
                if kind == "para":
                    _add_paragraph(doc, p, body, first.width, WD_ALIGN_PARAGRAPH, Pt, RGBColor)
                elif kind == "table":
                    rows = p.extract()
                    if not rows:
                        continue
                    cols = max(len(r) for r in rows)
                    tbl = doc.add_table(rows=len(rows), cols=cols)
                    tbl.style = "Table Grid"
                    for r_i, row in enumerate(rows):
                        for c_i in range(cols):
                            txt = (row[c_i] if c_i < len(row) and row[c_i] else "").strip()
                            cell = tbl.cell(r_i, c_i)
                            cell.text = txt
                            if r_i == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                    doc.add_paragraph()
                else:
                    im, full = p
                    x0, y0, x1, y1 = max(im["x0"], 0), max(im["top"], 0), im["x1"], im["bottom"]
                    sc = 150 / 72
                    crop = full.crop((int(x0 * sc), int(y0 * sc), int(x1 * sc), int(y1 * sc)))
                    buf = io.BytesIO()
                    crop.save(buf, format="PNG")
                    buf.seek(0)
                    doc.add_picture(buf, width=Pt(min(x1 - x0, avail)))
    doc.save(str(dst))
    return dst


def _add_paragraph(doc, p, body, page_width, ALIGN, Pt, RGBColor):
    runs = [r for r in p["runs"] if r["text"]]
    if not runs:
        return
    text = "".join(r["text"] for r in runs).strip()
    ratio = p["size"] / body if body else 1
    style = None
    if len(text) <= 140 and ratio >= 1.2:
        style = "Title" if ratio >= 2.0 else "Heading 1" if ratio >= 1.6 else "Heading 2" if ratio >= 1.35 else "Heading 3"
    m = _BULLET.match(runs[0]["text"])
    if m and not style:
        style = "List Bullet"
        runs[0] = dict(runs[0], text=runs[0]["text"][m.end():])
    para = doc.add_paragraph(style=style)
    for r in runs:
        run = para.add_run(r["text"])
        if not style or style == "List Bullet":
            run.font.size = Pt(r["size"] or body)
        if r["font"] and not style:
            run.font.name = r["font"]
        run.bold = True if r["bold"] else None
        run.italic = True if r["italic"] else None
        if r["color"] != "000000" and not style:
            run.font.color.rgb = RGBColor.from_string(r["color"])
    centre = abs((p["x0"] + p["x1"]) / 2 - page_width / 2) < 8 and p["x0"] > 90
    if centre:
        para.alignment = ALIGN.CENTER
    elif p["x0"] - 54 > 20 and not style:
        para.paragraph_format.left_indent = Pt(p["x0"] - 54)
    para.paragraph_format.space_after = Pt(6)
