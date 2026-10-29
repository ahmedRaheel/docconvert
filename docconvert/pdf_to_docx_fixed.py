"""PDF -> Word, exact-layout mode (for forms, letterheads, certificates, anything positional).

Flow-based conversion rebuilds a document as paragraphs and tables, which breaks
forms: letters typed into comb boxes drift out of their cells, rotated side labels
are mangled and fields overlap. This mode keeps every page exactly where it was:

* all lines, boxes, logos, check marks and rotated text stay in a page-sized
  picture placed *behind* the text (the text itself is cut out of that picture);
* every piece of real horizontal text becomes an editable Word text frame that is
  pinned to the same coordinates, with the original font, size, colour, bold and italic;
* letters typed into comb boxes get one frame per letter so they stay inside their cells.

Pure pdfplumber + pypdfium2: no extra packages beyond the core install.
"""
from __future__ import annotations

from .utils import clean_font, font_is_bold, font_is_italic, parse_pages

from .pdf_layout import looks_like_form, render_page  # noqa: F401  (looks_like_form re-exported)


def _frame(doc, s, page_w, dy):
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    x0, y0, x1, y1 = s["bbox"]
    h = max(y1 - y0, s["size"] * 1.1)
    w = min((x1 - x0) * 1.3 + 8, max(page_w - x0, x1 - x0 + 2))
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.left_indent = pf.right_indent = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(h)
    fr = OxmlElement("w:framePr")
    for k, v in (("w", int(w * 20)), ("h", int(h * 20) + 20), ("hRule", "exact"), ("wrap", "around"),
                 ("vAnchor", "page"), ("hAnchor", "page"), ("x", int(x0 * 20)), ("y", int((y0 + dy) * 20))):
        fr.set(qn("w:" + k), str(v))
    p._p.get_or_add_pPr().insert(0, fr)
    run = p.add_run(s["text"])
    f = run.font
    f.size = Pt(s["size"])
    name = clean_font(s["font"])
    if name:
        f.name = name
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        rf.set(qn("w:cs"), name)
        rf.set(qn("w:eastAsia"), name)
    f.bold = True if (font_is_bold(s["font"]) or s["flags"] & 16) else None
    f.italic = True if (font_is_italic(s["font"]) or s["flags"] & 2) else None
    f.color.rgb = RGBColor.from_string("%06X" % (s["color"] & 0xFFFFFF))


def _background(doc, png: bytes, w_pt, h_pt, first: bool):
    """Page-sized picture anchored to the page, behind the text."""
    import io

    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Emu, Pt

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(1)
    pf.page_break_before = not first
    run = p.add_run()
    run.font.size = Pt(1)
    run.add_picture(io.BytesIO(png), width=Emu(int(w_pt * 12700)), height=Emu(int(h_pt * 12700)))
    inline = run._r.xpath(".//wp:inline")[0]
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>')
    ext = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    anchor.insert(2, ext)
    anchor.insert(4, inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"))
    anchor.append(inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cNvGraphicFramePr"))
    anchor.append(inline.find("{http://schemas.openxmlformats.org/drawingml/2006/main}graphic"))
    inline.getparent().replace(inline, anchor)


def pdf_to_docx_fixed(src, dst, *, pages=None, dpi: int = 200, y_offset: float = 0.0, **_):
    import pdfplumber
    from docx import Document
    from docx.shared import Pt

    with pdfplumber.open(str(src)) as pl:
        idx = parse_pages(pages, len(pl.pages))
        fw, fh = float(pl.pages[idx[0]].width), float(pl.pages[idx[0]].height)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(fw), Pt(fh)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(0)
    sec.header_distance = sec.footer_distance = Pt(0)
    for para in list(doc.paragraphs):                       # drop the template's empty paragraph
        para._p.getparent().remove(para._p)

    for n, i in enumerate(idx):
        spans, png, (w, h) = render_page(src, i, dpi)
        for s in spans:
            _frame(doc, s, w, y_offset)
        _background(doc, png, w, h, first=(n == 0))
    doc.save(str(dst))
    return dst
