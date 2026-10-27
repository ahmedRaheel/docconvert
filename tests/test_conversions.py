import itertools
from pathlib import Path

import pytest

import docconvert
from docconvert import convert

FORMATS = ["pdf", "docx", "xlsx", "pptx"]


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    d = tmp_path_factory.mktemp("samples")
    # DOCX
    from docx import Document
    doc = Document()
    doc.add_heading("Quarterly Report", 1)
    p = doc.add_paragraph("Revenue grew ")
    p.add_run("strongly").bold = True
    p.add_run(" this quarter.")
    doc.add_paragraph("First point", style="List Bullet")
    t = doc.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([["Region", "Q1", "Q2"], ["North", "1,200", "1,500"], ["South", "900", "1,100"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    doc.save(d / "s.docx")
    # XLSX
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Sales"
    ws.append(["Item", "Qty", "Price"]); ws.append(["Widget", 4, 2.5]); ws.append(["Gadget", 10, 7.25])
    wb.save(d / "s.xlsx")
    # PPTX
    from pptx import Presentation
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"; s.placeholders[1].text = "Launch in Q3"
    prs.save(d / "s.pptx")
    # PDF (from the docx via the builtin engine so tests don't need LibreOffice)
    convert(d / "s.docx", d / "s.pdf", engine="builtin")
    return d


def _text(path: Path) -> str:
    ext = path.suffix
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return " ".join((p.extract_text() or "") for p in pdf.pages)
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        return " ".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    if ext == ".xlsx":
        from openpyxl import load_workbook
        return " ".join(str(c.value) for ws in load_workbook(path) for r in ws.iter_rows() for c in r if c.value is not None)
    from pptx import Presentation
    return " ".join(sh.text_frame.text for s in Presentation(path).slides for sh in s.shapes if sh.has_text_frame)


@pytest.mark.parametrize("src,dst", [p for p in itertools.permutations(FORMATS, 2)])
@pytest.mark.parametrize("engine", ["builtin", "auto"])
def test_every_pair(samples, tmp_path, src, dst, engine):
    out = convert(samples / f"s.{src}", tmp_path / f"out.{dst}", engine=engine)
    assert out.exists() and out.stat().st_size > 0
    assert _text(out).strip(), f"{src}->{dst} produced no text"


def test_pdf_to_xlsx_tables_are_numeric(samples, tmp_path):
    out = convert(samples / "s.pdf", tmp_path / "t.xlsx")
    from openpyxl import load_workbook
    vals = [c.value for ws in load_workbook(out) for r in ws.iter_rows() for c in r]
    assert 1200 in vals and 1500 in vals


def test_pdf_to_pptx_modes(samples, tmp_path):
    from pptx import Presentation
    ed = Presentation(convert(samples / "s.pdf", tmp_path / "e.pptx", mode="editable"))
    assert any(sh.has_text_frame and "Revenue" in sh.text_frame.text for sh in ed.slides[0].shapes)
    im = Presentation(convert(samples / "s.pdf", tmp_path / "i.pptx", mode="image"))
    assert any(sh.shape_type == 13 for sh in im.slides[0].shapes)


def test_page_selection_and_errors(samples, tmp_path):
    convert(samples / "s.pdf", tmp_path / "p.docx", pages="1")
    with pytest.raises(docconvert.UnsupportedConversionError):
        convert(samples / "s.pdf", tmp_path / "x.pdf")
    with pytest.raises(docconvert.UnsupportedConversionError):
        convert(samples / "s.pdf", tmp_path / "x.txt")
    with pytest.raises(FileNotFoundError):
        convert(tmp_path / "nope.pdf", to="docx")


HTML = """<html><body><h1>Invoice 42</h1><p>Hello <b>bold</b> world.</p>
<ul><li>One</li><li>Two</li></ul>
<table><tr><th>Item</th><th>Cost</th></tr><tr><td>Pen</td><td>1,500</td></tr></table></body></html>"""


@pytest.mark.parametrize("engine", ["builtin", "auto"])
@pytest.mark.parametrize("target", ["pdf", "docx", "xlsx", "pptx"])
def test_html_source(tmp_path, engine, target):
    src = tmp_path / "p.html"
    src.write_text(HTML, encoding="utf-8")
    out = convert(src, tmp_path / f"p.{target}", engine=engine)
    assert "Invoice" in _text(out) or target == "xlsx"
    assert _text(out).strip()


def test_html_string(tmp_path):
    out = docconvert.html_string_to_pdf("<h1>Hi there</h1>", tmp_path / "s.pdf", engine="builtin")
    assert "Hi there" in _text(out)
