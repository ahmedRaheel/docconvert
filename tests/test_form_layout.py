"""Exact-copy mode for forms (PDF -> Word / PowerPoint). No PyMuPDF needed."""
import zipfile

import pytest

import docconvert


@pytest.fixture()
def form_pdf(tmp_path):
    """Small form: label, comb boxes, text hidden under a white box, rotated side label."""
    from reportlab.pdfgen import canvas
    out = tmp_path / "form.pdf"
    c = canvas.Canvas(str(out), pagesize=(300, 300))
    c.setFont("Helvetica", 9)
    c.drawString(20, 270, "Name of Candidate")
    for i, ch in enumerate("OWAIS"):
        x = 120 + i * 14
        c.rect(x, 266, 14, 14)
        c.drawString(x + 4, 270, ch)
    c.drawString(20, 220, "HIDDENOLD")
    c.setFillColorRGB(1, 1, 1); c.rect(15, 212, 110, 17, stroke=0, fill=1); c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 10)
    c.drawString(20, 220, "Session 2026")
    c.saveState(); c.translate(285, 60); c.rotate(90); c.setFont("Helvetica", 9); c.drawString(0, 0, "Side label"); c.restoreState()
    c.save()
    return out


def _xml(docx):
    with zipfile.ZipFile(docx) as z:
        return z.read("word/document.xml").decode("utf8")


def test_fixed_layout_pins_text_and_keeps_boxes(form_pdf, tmp_path):
    xml = _xml(docconvert.convert(form_pdf, tmp_path / "f.docx", layout="fixed"))
    assert xml.count("<w:framePr") >= 6                     # label + one frame per comb letter + session
    assert "Name of Candidate" in xml and "Session 2026" in xml
    assert "HIDDENOLD" not in xml                           # text under a white box is not resurrected
    assert "Side label" not in xml                          # rotated text stays in the picture
    assert 'behindDoc="1"' in xml                           # lines/boxes ride along as a background picture


def test_flow_layout_has_no_frames(form_pdf, tmp_path):
    out = docconvert.convert(form_pdf, tmp_path / "flow.docx", layout="flow", engine="builtin")
    assert "<w:framePr" not in _xml(out)


def test_pdf_form_to_pptx_exact_mode(form_pdf, tmp_path):
    from pptx import Presentation
    prs = Presentation(docconvert.convert(form_pdf, tmp_path / "f.pptx", mode="exact"))
    shapes = list(prs.slides[0].shapes)
    assert sum(sh.shape_type == 13 for sh in shapes) == 1                       # page graphics as one picture
    texts = [sh.text_frame.text for sh in shapes if sh.has_text_frame]
    assert "Name of Candidate" in texts and "Session 2026" in texts
    assert {"O", "W", "A", "I", "S"} <= set(texts)                              # one box per comb letter
    assert "HIDDENOLD" not in texts


def test_exact_mode_does_not_need_pymupdf(form_pdf, tmp_path, monkeypatch):
    import sys
    monkeypatch.setitem(sys.modules, "fitz", None)
    monkeypatch.setitem(sys.modules, "pymupdf", None)
    assert docconvert.convert(form_pdf, tmp_path / "x.docx", layout="fixed").exists()


def test_looks_like_form_detects_comb_boxes(tmp_path):
    from reportlab.pdfgen import canvas
    out = tmp_path / "comb.pdf"
    c = canvas.Canvas(str(out), pagesize=(400, 200))
    for row in range(3):
        for i, ch in enumerate("ABCDEFGHIJKL"):
            x = 40 + i * 20
            c.rect(x, 150 - row * 30, 20, 20)
            c.drawString(x + 6, 157 - row * 30, ch)
    c.save()
    assert docconvert.pdf_layout.looks_like_form(out) is True
