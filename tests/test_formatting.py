"""Formatting (font size, style, colour, fills, alignment...) must survive every conversion."""
import datetime as dt

import pdfplumber
import pytest

import docconvert as dc


def _chars(pdf):
    with pdfplumber.open(pdf) as p:
        return [c for pg in p.pages for c in pg.chars]


def _find(chars, text):
    s = "".join(c["text"] for c in chars)
    i = s.index(text)
    return chars[i:i + len(text)]


def rgb(c):
    return tuple(round(x, 2) for x in (c["non_stroking_color"] or ()))


@pytest.fixture(scope="module")
def styled_docx(tmp_path_factory):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    d = Document()
    h = d.add_heading("Report Title", 1)
    h.runs[0].font.color.rgb = RGBColor(0xC0, 0x39, 0x2B); h.runs[0].font.size = Pt(26)
    p = d.add_paragraph()
    r = p.add_run("Red14"); r.font.size = Pt(14); r.bold = True; r.italic = True; r.font.color.rgb = RGBColor(0xCC, 0, 0)
    r = p.add_run(" green"); r.underline = True; r.font.color.rgb = RGBColor(0, 0x80, 0)
    q = d.add_paragraph("Centered9"); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.runs[0].font.size = Pt(9)
    r = d.add_paragraph().add_run("TimesBlue"); r.font.name = "Times New Roman"; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0, 0, 0xFF)
    path = tmp_path_factory.mktemp("f") / "s.docx"
    d.save(path)
    return path


def test_docx_to_pdf_keeps_size_colour_style_font(styled_docx, tmp_path):
    ch = _chars(dc.convert(styled_docx, tmp_path / "o.pdf", engine="builtin"))
    t = _find(ch, "Report Title")
    assert round(t[0]["size"]) == 26 and rgb(t[0]) == (0.75, 0.22, 0.17)
    red = _find(ch, "Red14")
    assert round(red[0]["size"]) == 14 and "Bold" in red[0]["fontname"] and "Italic" in red[0]["fontname"] and rgb(red[0]) == (0.8, 0.0, 0.0)
    assert round(_find(ch, "Centered9")[0]["size"]) == 9
    tb = _find(ch, "TimesBlue")
    assert "Serif" in tb[0]["fontname"] and rgb(tb[0]) == (0.0, 0.0, 1.0)
    # centred paragraph sits in the middle of the page
    page_mid = 595 / 2 if False else None
    cx = (_find(ch, "Centered9")[0]["x0"] + _find(ch, "Centered9")[-1]["x1"]) / 2
    assert 250 < cx < 360


def test_xlsx_to_pdf_keeps_fonts_fills_formats(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active
    ws["A1"] = "Title"; ws["A1"].font = Font(size=18, bold=True, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws.merge_cells("A1:C1"); ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = 15234.5; ws["A2"].number_format = '"$"#,##0.00'
    ws["B2"] = 0.125; ws["B2"].number_format = "0.0%"
    ws["C2"] = dt.date(2026, 3, 5); ws["C2"].number_format = "dd mmm yyyy"
    ws["A3"] = "small italic"; ws["A3"].font = Font(size=8, italic=True, color="008000")
    wb.save(tmp_path / "s.xlsx")
    pdf = dc.convert(tmp_path / "s.xlsx", tmp_path / "o.pdf", engine="builtin")
    ch = _chars(pdf)
    text = "".join(c["text"] for c in ch)
    assert "$15,234.50" in text and "12.5%" in text and "05 Mar 2026" in text
    t = _find(ch, "Title")
    assert round(t[0]["size"]) == 18 and "Bold" in t[0]["fontname"] and rgb(t[0]) == (1.0, 1.0, 1.0)
    s = _find(ch, "small italic")
    assert round(s[0]["size"]) == 8 and "Italic" in s[0]["fontname"] and rgb(s[0]) == (0.0, 0.5, 0.0)
    with pdfplumber.open(pdf) as p:
        assert any(r.get("fill") and tuple(round(x, 2) for x in r["non_stroking_color"]) == (0.12, 0.31, 0.47) for r in p.pages[0].rects)


def test_pptx_to_pdf_resolves_inherited_and_explicit_formatting(tmp_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Roadmap"
    tf = s.placeholders[1].text_frame; tf.text = "Inherited"
    r = tf.add_paragraph().add_run(); r.text = "Explicit"; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
    prs.save(tmp_path / "s.pptx")
    ch = _chars(dc.convert(tmp_path / "s.pptx", tmp_path / "o.pdf", engine="builtin"))
    assert round(_find(ch, "Roadmap")[0]["size"]) == 44           # from the master title style
    assert round(_find(ch, "Inherited")[0]["size"]) == 32         # from the master body style
    e = _find(ch, "Explicit")
    assert round(e[0]["size"]) == 24 and "Bold" in e[0]["fontname"] and rgb(e[0]) == (0.75, 0.0, 0.0)


def test_pdf_to_word_keeps_run_formatting(styled_docx, tmp_path):
    pdf = dc.convert(styled_docx, tmp_path / "o.pdf", engine="builtin")
    out = dc.convert(pdf, tmp_path / "b.docx", engine="builtin", layout="flow")
    from docx import Document
    runs = [r for p in Document(out).paragraphs for r in p.runs]
    title = next(r for r in runs if "Report Title" in r.text)
    assert round(title.font.size.pt) == 26 and str(title.font.color.rgb) == "C0392B"
    red = next(r for r in runs if "Red14" in r.text)
    assert round(red.font.size.pt) == 14 and red.bold and red.italic and str(red.font.color.rgb) == "CC0000"


def test_pdf_to_pptx_keeps_text_style_and_rules(styled_docx, tmp_path):
    pdf = dc.convert(styled_docx, tmp_path / "o.pdf", engine="builtin")
    from pptx import Presentation
    prs = Presentation(dc.convert(pdf, tmp_path / "p.pptx", mode="editable"))
    runs = [r for sh in prs.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    red = next(r for r in runs if "Red14" in r.text)
    assert round(red.font.size.pt) == 14 and red.font.bold and red.font.italic and str(red.font.color.rgb) == "CC0000"


def test_pdf_to_xlsx_copies_cell_formatting(tmp_path):
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    data = [["Quarterly", "", ""], ["Region", "Units", "Growth"], ["North", "1,200", "12.5%"], ["South", "980", "-4.0%"]]
    t = Table(data, colWidths=[110, 80, 80])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black), ("SPAN", (0, 0), (-1, 0)),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, 0), 16), ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#C00000")), ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#DDEBF7")),
        ("ALIGN", (1, 2), (-1, -1), "RIGHT")]))
    SimpleDocTemplate(str(tmp_path / "t.pdf")).build([t])
    from openpyxl import load_workbook
    ws = load_workbook(dc.convert(tmp_path / "t.pdf", tmp_path / "t.xlsx")).worksheets[0]
    title = ws["A1"]
    assert title.value == "Quarterly" and round(title.font.sz) == 16 and title.font.b and title.font.color.rgb == "FFFFFFFF"
    assert title.fill.fgColor.rgb == "FF1F4E79" and title.alignment.horizontal == "center"
    assert "A1:C1" in {str(r) for r in ws.merged_cells.ranges}
    hdr = ws["A2"]
    assert hdr.font.color.rgb == "FFC00000" and hdr.font.b and hdr.fill.fgColor.rgb == "FFDDEBF7"
    assert ws["B3"].value == 1200 and ws["B3"].number_format == "#,##0" and ws["B3"].alignment.horizontal == "right"
    assert ws["C3"].value == 0.125 and ws["C3"].number_format == "0.0%" or ws["C3"].number_format.endswith("%")
    assert ws["A3"].border.left.style is not None


def test_html_css_is_applied_by_builtin_renderer(tmp_path):
    (tmp_path / "s.html").write_text("""<html><head><style>
      h1 { color:#1F4E79; text-align:center } .warn { color:#C00000; font-weight:bold; font-size:14pt }
      th { background:#1F4E79; color:white }</style></head><body>
      <h1>Title</h1><p class="warn">Alert</p>
      <p style="color:rgb(0,128,0);font-size:16px;text-align:right">Green</p>
      <table border="1"><tr><th>H</th></tr><tr><td style="background:#E2EFDA">v</td></tr></table></body></html>""")
    pdf = dc.convert(tmp_path / "s.html", tmp_path / "o.pdf", engine="builtin")
    ch = _chars(pdf)
    t = _find(ch, "Title")
    assert round(t[0]["size"]) == 24 and rgb(t[0]) == (0.12, 0.31, 0.47)
    a = _find(ch, "Alert")
    assert round(a[0]["size"]) == 14 and "Bold" in a[0]["fontname"] and rgb(a[0]) == (0.75, 0.0, 0.0)
    g = _find(ch, "Green")
    assert round(g[0]["size"]) == 12 and rgb(g[0]) == (0.0, 0.5, 0.0) and g[0]["x0"] > 300       # right aligned
    with pdfplumber.open(pdf) as p:
        fills = {tuple(round(x, 2) for x in r["non_stroking_color"]) for r in p.pages[0].rects if r.get("fill")}
    assert (0.12, 0.31, 0.47) in fills and (0.89, 0.94, 0.85) in fills
