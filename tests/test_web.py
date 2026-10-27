import io
import re

import pytest

pytest.importorskip("flask")
from web.app import app  # noqa: E402


@pytest.fixture()
def client():
    app.config["TESTING"] = True
    return app.test_client()


def _docx_bytes():
    from docx import Document
    d = Document(); d.add_heading("Hello web", 1); d.add_paragraph("Body text")
    b = io.BytesIO(); d.save(b); return b.getvalue()


def test_index_and_health(client):
    r = client.get("/")
    assert r.status_code == 200 and b"Turn" in r.data and b'id="boot"' in r.data
    assert client.get("/healthz").json == {"status": "ok"}


def test_convert_download_and_zip(client):
    ids = []
    for _ in range(2):
        r = client.post("/api/convert", data={"file": (io.BytesIO(_docx_bytes()), "My Report.docx"), "to": "pdf",
                                              "options": '{"engine":"builtin"}'}, content_type="multipart/form-data")
        assert r.status_code == 200, r.json
        assert r.json["name"] == "My_Report.pdf"
        ids.append(r.json["id"])
    d = client.get(f"/api/download/{ids[0]}")
    assert d.status_code == 200 and d.data.startswith(b"%PDF")
    z = client.get("/api/download-all?ids=" + ",".join(ids))
    assert z.status_code == 200 and z.data[:2] == b"PK"


@pytest.mark.parametrize("data,status", [
    ({"to": "pdf"}, 400),
    ({"file": (io.BytesIO(b"x"), "a.exe"), "to": "pdf"}, 415),
    ({"file": (io.BytesIO(b"x"), "a.docx"), "to": "docx"}, 400),
    ({"file": (io.BytesIO(b"x"), "a.docx"), "to": "pdf", "options": '{"pages":"1;rm -rf"}'}, 400),
])
def test_rejections(client, data, status):
    assert client.post("/api/convert", data=data, content_type="multipart/form-data").status_code == status


def test_bad_file_gives_readable_error(client):
    r = client.post("/api/convert", data={"file": (io.BytesIO(b"not a pdf"), "bad.pdf"), "to": "docx"},
                    content_type="multipart/form-data")
    assert r.status_code in (422, 500) and "error" in r.json


def test_download_rejects_traversal(client):
    assert client.get("/api/download/..%2f..%2fetc").status_code == 404
    assert client.get("/api/download/" + "0" * 32).status_code == 404


def test_basic_auth(client, monkeypatch):
    monkeypatch.setenv("BASIC_AUTH_USER", "u"); monkeypatch.setenv("BASIC_AUTH_PASSWORD", "p")
    assert client.get("/").status_code == 401
    assert client.get("/healthz").status_code == 200
    import base64
    h = {"Authorization": "Basic " + base64.b64encode(b"u:p").decode()}
    assert client.get("/", headers=h).status_code == 200


def test_prefix_support(client):
    r = client.get("/", headers={"X-Forwarded-Prefix": "/convert"})
    assert b'/convert/static/css/app.css' in r.data
